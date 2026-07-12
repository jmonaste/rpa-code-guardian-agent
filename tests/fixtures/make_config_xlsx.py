"""Regenerate tests/fixtures/sample-project/Data/Config.xlsx.

Run once when the fixture needs to change:
    python tests/fixtures/make_config_xlsx.py
"""

from pathlib import Path

from openpyxl import Workbook

OUT = Path(__file__).parent / "sample-project" / "Data" / "Config.xlsx"

wb = Workbook()

settings = wb.active
settings.title = "Settings"
settings.append(["Name", "Value", "Description"])
settings.append(["OrchestratorQueueName", "ACME_Invoices", "Queue holding the invoice transactions"])
settings.append(["InvoiceSystemUrl", "https://invoicing.acme.example", "URL of the invoicing web application"])
settings.append(["ReportPath", "\\\\fileserver\\reports", "Unused legacy report folder"])
settings.append(["ApplicationPassword", "SuperSecret123", "Credential that must never leak"])

constants = wb.create_sheet("Constants")
constants.append(["Name", "Value", "Description"])
constants.append(["MaxRetryNumber", "2", "Retries per transaction"])
constants.append(["logF_BusinessProcessName", "ACME-InvoiceProcessing", ""])

assets = wb.create_sheet("Assets")
assets.append(["Name", "Asset"])
assets.append(["InvoiceSystemCredential", "ACME_Invoice_Credential"])

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"written {OUT}")
