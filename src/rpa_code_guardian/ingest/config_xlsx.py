"""Read the REFramework ``Config.xlsx`` workbook (Settings / Constants / Assets)."""

from __future__ import annotations

import re
from pathlib import Path

from ..model.ir import ConfigEntry

_SECRET_RE = re.compile(r"password|secret|token|api[_-]?key|credential", re.IGNORECASE)


def parse_config_xlsx(path: Path) -> list[ConfigEntry]:
    """Extract Name/Value/Description rows from every sheet of the workbook.

    Values whose *name* looks secret are redacted so they never reach the LLM
    or the generated document. Returns [] if the file cannot be read.
    """
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return []

    entries: list[ConfigEntry] = []
    try:
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                continue
            cols = {str(c).strip().lower(): i for i, c in enumerate(header) if c is not None}
            name_i = cols.get("name", cols.get("asset", 0))
            value_i = cols.get("value", cols.get("asset", 1))
            desc_i = cols.get("description")
            for row in rows:
                if row is None or name_i >= len(row) or row[name_i] is None:
                    continue
                name = str(row[name_i]).strip()
                if not name:
                    continue
                value = str(row[value_i]).strip() if value_i < len(row) and row[value_i] is not None else ""
                if _SECRET_RE.search(name):
                    value = "(redacted)"
                desc = ""
                if desc_i is not None and desc_i < len(row) and row[desc_i] is not None:
                    desc = str(row[desc_i]).strip()
                entries.append(
                    ConfigEntry(sheet=ws.title, name=name, value=value[:200], description=desc[:200])
                )
    finally:
        wb.close()
    return entries
